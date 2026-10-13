from __future__ import annotations

import base64
import io
import json
import os
import re
import time
from pathlib import Path
from typing import Any

import streamlit as st
import streamlit.components.v1 as components
from docx import Document
from google import genai
from google.genai import types
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pypdf import PdfReader

APP_DIR = Path(__file__).parent
OUTPUT_DIR = APP_DIR / "generated"
OUTPUT_DIR.mkdir(exist_ok=True)

st.set_page_config(page_title="Prism AI", page_icon="✨", layout="wide")

# API key is read from Streamlit secrets first, then an environment variable.
def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""
    return key or os.getenv("GEMINI_API_KEY", "")


def get_client() -> genai.Client:
    key = get_api_key()
    if not key:
        raise RuntimeError("GEMINI_API_KEY সেট করা নেই। README-এর setup ধাপ দেখো।")
    return genai.Client(api_key=key)


def extract_document(uploaded_file) -> str:
    """Extract readable text from PDF, DOCX or TXT files."""
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            return "[এই PDF থেকে selectable text পাওয়া যায়নি। সম্ভবত এটি scanned PDF; OCR যোগ করতে হবে।]"
        return text
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if name.endswith(".txt"):
        return raw.decode("utf-8", errors="replace")
    raise ValueError("শুধু PDF, DOCX অথবা TXT ফাইল দাও।")


def ask_gemini(prompt: str, image: Image.Image | None = None, history: list[dict] | None = None) -> str:
    client = get_client()
    model = os.getenv("GEMINI_TEXT_MODEL", "gemini-3.8-flash")
    contents: list[Any] = []
    if history:
        # Keep a short conversation context so requests remain inexpensive.
        for item in history[-8:]:
            contents.append(f"{item['role'].upper()}: {item['content']}")
    contents.append(prompt)
    if image is not None:
        contents.append(image)
    response = client.models.generate_content(model=model, contents=contents)
    return response.text or "দুঃখিত, এইবার কোনো text response পাওয়া যায়নি।"


def save_generated_image(prompt: str) -> Path:
    client = get_client()
    model = os.getenv("GEMINI_IMAGE_MODEL", "gemini-3.1-flash-image")
    response = client.models.generate_content(model=model, contents=prompt)
    for part in response.parts:
        if getattr(part, "inline_data", None) is not None:
            image = part.as_image()
            path = OUTPUT_DIR / f"image_{int(time.time())}.png"
            image.save(path)
            return path
    raise RuntimeError("মডেল কোনো ছবি ফেরত দেয়নি। তোমার API key/model access পরীক্ষা করো।")


def create_presentation(topic: str, language: str, slide_count: int) -> Path:
    prompt = f"""Create a PowerPoint outline for topic: {topic!r}. Language: {language}. Slides: {slide_count}.
Return ONLY valid JSON array. Each item: {{"title":"short title","bullets":["point 1","point 2","point 3"]}}.
No markdown fences. Keep bullets concise and factual. First slide is an introduction and final slide is a conclusion."""
    raw = ask_gemini(prompt)
    match = re.search(r"\[.*\]", raw, flags=re.S)
    if not match:
        raise RuntimeError("AI সঠিক slide outline তৈরি করতে পারেনি। আবার চেষ্টা করো।")
    slides = json.loads(match.group(0))
    prs = Presentation()
    for idx, item in enumerate(slides[:slide_count]):
        layout = prs.slide_layouts[0] if idx == 0 else prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = str(item.get("title", f"Slide {idx + 1}"))
        if idx == 0 and len(slide.placeholders) < 2:
            continue
        if len(slide.placeholders) > 1:
            body = slide.placeholders[1].text_frame
            body.clear()
            bullets = item.get("bullets", [])
            if not bullets:
                bullets = [topic]
            for bidx, bullet in enumerate(bullets):
                p = body.paragraphs[0] if bidx == 0 else body.add_paragraph()
                p.text = str(bullet)
                p.level = 0
                p.font.size = Pt(22)
    path = OUTPUT_DIR / f"slides_{int(time.time())}.pptx"
    prs.save(path)
    return path


def create_video(prompt: str) -> Path:
    """Optional Veo integration. Availability and billing depend on the API account."""
    client = get_client()
    model = os.getenv("GEMINI_VIDEO_MODEL", "veo-3.1-generate-preview")
    operation = client.models.generate_videos(model=model, prompt=prompt)
    status = st.empty()
    while not operation.done:
        status.info("ভিডিও তৈরি হচ্ছে... এটি কয়েক মিনিট লাগতে পারে।")
        time.sleep(10)
        operation = client.operations.get(operation)
    status.empty()
    generated = getattr(operation, "response", None)
    if generated is None:
        raise RuntimeError("ভিডিও তৈরি হয়নি। API access, quota এবং billing পরীক্ষা করো।")
    videos = getattr(generated, "generated_videos", None) or []
    if not videos:
        raise RuntimeError("API কোনো video file ফেরত দেয়নি।")
    video = videos[0].video
    path = OUTPUT_DIR / f"video_{int(time.time())}.mp4"
    client.files.download(file=video, download_path=str(path))
    return path


def speak_text(text: str, language: str) -> None:
    # Uses the browser's built-in speech synthesis; no paid TTS API required.
    safe_text = json.dumps(text)
    lang = "bn-BD" if language == "বাংলা" else "en-US"
    components.html(
        f"""<button id="speak-btn" style="border:0;border-radius:10px;padding:8px 12px;cursor:pointer;background:#7c3aed;color:white;font-weight:600">🔊 উত্তর শুনুন / Read aloud</button>
        <script>
        const btn = document.getElementById('speak-btn');
        btn.addEventListener('click', () => {{
          const utterance = new SpeechSynthesisUtterance({safe_text});
          utterance.lang = {json.dumps(lang)};
          window.speechSynthesis.cancel();
          window.speechSynthesis.speak(utterance);
        }});
        </script>""",
        height=48,
    )


# ---- UI ----
st.markdown("""<style>
.block-container {max-width: 1200px; padding-top: 1.8rem;}
.hero {padding: 1.4rem 1.6rem; border-radius: 20px; background: linear-gradient(120deg,#172554,#312e81,#0f766e); color:white; margin-bottom:1rem;}
.hero h1 {margin:0; color:white; font-size:2.2rem;}
.hero p {margin:.4rem 0 0 0; color:#e0e7ff;}
</style>""", unsafe_allow_html=True)
st.markdown('<div class="hero"><h1>✨ Prism AI</h1><p>Your bilingual multimodal assistant · বাংলা ও English</p></div>', unsafe_allow_html=True)

with st.sidebar:
    st.header("⚙️ Settings")
    ui_language = st.selectbox("উত্তরের ভাষা / Response language", ["বাংলা", "English", "প্রশ্ন অনুযায়ী / Auto"])
    st.caption("API key code-এর মধ্যে লিখবে না। .streamlit/secrets.toml-এ রাখবে।")
    if get_api_key():
        st.success("Gemini API key পাওয়া গেছে")
    else:
        st.warning("Gemini API key সেট করা হয়নি")
    st.markdown("---")
    st.caption("Free usage limits এবং feature access Google account/region অনুযায়ী আলাদা হতে পারে।")

chat_tab, doc_tab, voice_tab, image_tab, slides_tab, video_tab = st.tabs([
    "💬 Chat & Image", "📄 Documents", "🎙️ Voice", "🎨 Image Generator", "📊 Slides", "🎬 Video"
])

if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

with chat_tab:
    st.subheader("AI Chat & Image Understanding")
    uploaded_image = st.file_uploader("ছবি upload করো (optional)", type=["png", "jpg", "jpeg", "webp"], key="chat_image")
    if uploaded_image:
        st.image(uploaded_image, caption="তোমার আপলোড করা ছবি", use_container_width=True)
    for item in st.session_state.chat_history:
        with st.chat_message(item["role"]):
            st.markdown(item["content"])
    question = st.chat_input("এখানে প্রশ্ন লেখো... / Ask Prism AI")
    if question:
        st.session_state.chat_history.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)
        try:
            img = Image.open(uploaded_image).convert("RGB") if uploaded_image else None
            language_instruction = "Answer in Bangla." if ui_language == "বাংলা" else "Answer in English." if ui_language == "English" else "Answer in the same language as the user."
            answer = ask_gemini(f"{language_instruction}\nUser question: {question}", img, st.session_state.chat_history[:-1])
            st.session_state.chat_history.append({"role": "assistant", "content": answer})
            with st.chat_message("assistant"):
                st.markdown(answer)
                speak_text(answer, "বাংলা" if ui_language == "বাংলা" else "English")
        except Exception as exc:
            st.error(f"সমস্যা হয়েছে: {exc}")
    if st.button("🧹 Chat history মুছো"):
        st.session_state.chat_history = []
        st.rerun()

with doc_tab:
    st.subheader("PDF / DOCX / TXT বুঝে নেওয়া")
    doc_file = st.file_uploader("Document upload করো", type=["pdf", "docx", "txt"], key="document_file")
    if doc_file:
        st.caption(f"ফাইল: {doc_file.name} · {doc_file.size:,} bytes")
        if st.button("📖 Document বিশ্লেষণ করো", type="primary"):
            try:
                with st.spinner("ফাইল পড়ছি..."):
                    extracted = extract_document(doc_file)
                st.session_state.document_text = extracted[:100000]
                st.success("Document থেকে text নেওয়া হয়েছে। নিচে প্রশ্ন করো।")
                st.text_area("Extracted text preview", st.session_state.document_text[:6000], height=180)
            except Exception as exc:
                st.error(f"ফাইল পড়তে সমস্যা: {exc}")
    if st.session_state.get("document_text"):
        doc_question = st.text_input("Document সম্পর্কে কী জানতে চাও?", placeholder="সারাংশ দাও / মূল পয়েন্টগুলো বলো")
        if st.button("উত্তর দাও", key="ask_doc") and doc_question:
            try:
                lang = "Bangla" if ui_language == "বাংলা" else "English" if ui_language == "English" else "same language as the question"
                answer = ask_gemini(f"Answer in {lang}. Use only the document when possible. If information is absent, say so.\nDOCUMENT:\n{st.session_state.document_text}\n\nQUESTION:\n{doc_question}")
                st.markdown(answer)
            except Exception as exc:
                st.error(f"সমস্যা হয়েছে: {exc}")

with voice_tab:
    st.subheader("Voice Input & Voice Response")
    st.write("Record button চাপো, কথা বলো, তারপর recording থেকে AI-কে প্রশ্ন করো।")
    audio_file = st.audio_input("🎙️ কথা record করো / Record your voice")
    if audio_file:
        st.audio(audio_file)
        if st.button("শুনে উত্তর দাও / Understand audio", type="primary"):
            try:
                client = get_client()
                audio_bytes = audio_file.getvalue()
                mime = audio_file.type or "audio/wav"
                language_instruction = "Reply in Bangla." if ui_language == "বাংলা" else "Reply in English." if ui_language == "English" else "Reply in the language spoken in the audio."
                response = client.models.generate_content(
                    model=os.getenv("GEMINI_TEXT_MODEL", "gemini-3.8-flash"),
                    contents=[types.Part.from_bytes(data=audio_bytes, mime_type=mime), language_instruction + " Understand the user's speech, answer their request, and briefly state what you understood."],
                )
                answer = response.text or "কথা বুঝতে পারিনি। আরেকবার চেষ্টা করো।"
                st.markdown(answer)
                speak_text(answer, "বাংলা" if ui_language == "বাংলা" else "English")
            except Exception as exc:
                st.error(f"Voice processing-এ সমস্যা: {exc}")
    st.info("Voice input browser-এর microphone permission চায়। Voice output তোমার browser/OS-এর speech voice ব্যবহার করে।")

with image_tab:
    st.subheader("AI Image Generator")
    image_prompt = st.text_area("কী ছবি বানাতে চাও?", placeholder="A cinematic illustration of a futuristic Dhaka skyline at sunset...", height=100)
    if st.button("✨ ছবি তৈরি করো", type="primary", key="generate_image"):
        if not image_prompt.strip():
            st.warning("আগে ছবির বর্ণনা লেখো।")
        else:
            try:
                with st.spinner("ছবি তৈরি হচ্ছে..."):
                    path = save_generated_image(image_prompt)
                st.image(str(path), caption="Generated by Prism AI", use_container_width=True)
                st.download_button("⬇️ ছবি Download", path.read_bytes(), file_name=path.name, mime="image/png")
            except Exception as exc:
                st.error(f"ছবি তৈরি হয়নি: {exc}")
                st.caption("Image-generation model-এর API access ও quota পরীক্ষা করো।")

with slides_tab:
    st.subheader("PowerPoint Slide Generator")
    topic = st.text_input("Presentation-এর বিষয়", placeholder="Introduction to Artificial Intelligence")
    col1, col2 = st.columns(2)
    with col1:
        slide_language = st.selectbox("Slide language", ["Bangla", "English"], key="slide_language")
    with col2:
        slide_count = st.slider("কয়টি slide?", 3, 12, 6)
    if st.button("📊 PowerPoint তৈরি করো", type="primary"):
        if not topic.strip():
            st.warning("Presentation-এর বিষয় লেখো।")
        else:
            try:
                with st.spinner("Slide content তৈরি হচ্ছে..."):
                    path = create_presentation(topic, slide_language, slide_count)
                st.success("PowerPoint তৈরি হয়েছে!")
                st.download_button("⬇️ Download .pptx", path.read_bytes(), file_name=path.name, mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")
            except Exception as exc:
                st.error(f"Slide তৈরি হয়নি: {exc}")

with video_tab:
    st.subheader("AI Video Generator — Optional")
    st.warning("ভিডিও generation-এর জন্য Veo API access লাগতে পারে এবং এটি paid হতে পারে। Generate করার আগে Google AI account-এর pricing ও quota দেখে নাও।")
    video_prompt = st.text_area("ভিডিওর বর্ণনা", placeholder="An 8-second cinematic shot of a river boat at sunrise in Bangladesh...")
    if st.button("🎬 ভিডিও তৈরি করো", type="primary"):
        if not video_prompt.strip():
            st.warning("আগে ভিডিওর বর্ণনা লেখো।")
        else:
            try:
                with st.spinner("ভিডিও তৈরি হচ্ছে; কয়েক মিনিট লাগতে পারে..."):
                    path = create_video(video_prompt)
                st.video(path.read_bytes())
                st.download_button("⬇️ Download video", path.read_bytes(), file_name=path.name, mime="video/mp4")
            except Exception as exc:
                st.error(f"ভিডিও তৈরি হয়নি: {exc}")
                st.caption("Veo model access, billing, quota এবং google-genai SDK version পরীক্ষা করো।")

st.markdown("---")
st.caption("Prism AI · Prototype · Do not upload confidential documents unless you understand the provider's data terms.")
